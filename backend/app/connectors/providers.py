"""Provider HTTP adapters; tokens and org persistence live in the control plane."""
import base64
import html
import re
import io
import zipfile
from pathlib import PurePosixPath
from urllib.parse import quote, urlparse, parse_qs
from .common import request, require, encoded, Item, Batch, ConnectorError, ProviderError

GOOGLE_SCOPES = {
    'gmail': 'https://www.googleapis.com/auth/gmail.readonly',
    'drive': 'https://www.googleapis.com/auth/drive.readonly',
}
RAMP_HOSTS = {'production': 'https://api.ramp.com', 'sandbox': 'https://demo-api.ramp.com'}
PLAID_HOSTS = {'production': 'https://production.plaid.com', 'sandbox': 'https://sandbox.plaid.com'}


def google_token(client, params):
    import os
    require('GOOGLE_CLIENT_ID', 'GOOGLE_CLIENT_SECRET')
    return request(client, 'POST', 'https://oauth2.googleapis.com/token', data={
        'client_id': os.environ['GOOGLE_CLIENT_ID'], 'client_secret': os.environ['GOOGLE_CLIENT_SECRET'], **params})


def ramp_token(client, credentials, environment):
    return request(client, 'POST', RAMP_HOSTS[environment] + '/developer/v1/token',
                   auth=(credentials['client_id'], credentials['client_secret']),
                   data={'grant_type': 'client_credentials', 'scope': 'bills:read transactions:read'})


def plaid(client, path, environment, **body):
    import os
    require('PLAID_CLIENT_ID', 'PLAID_SECRET')
    return request(client, 'POST', PLAID_HOSTS[environment] + path,
                   headers={'Plaid-Version': '2020-09-14'},
                   json={'client_id': os.environ['PLAID_CLIENT_ID'], 'secret': os.environ['PLAID_SECRET'], **body})


def decode64(value):
    return base64.urlsafe_b64decode(value + '=' * (-len(value) % 4))


def gmail_message(client, headers, mid):
    root = 'https://gmail.googleapis.com/gmail/v1/users/me/messages/' + quote(mid, safe='')
    message = request(client, 'GET', root, headers=headers, params={'format': 'full'})
    return root, message


def gmail_items(client, headers, mid, labels):
    root, message = gmail_message(client, headers, mid)
    if not set(labels).issubset(message.get('labelIds', [])):
        return [], True
    payload = message.get('payload', {})
    parts = [payload]
    text, attachments = [], []
    while parts:
        part = parts.pop()
        parts.extend(reversed(part.get('parts', [])))
        body = part.get('body', {})
        filename = part.get('filename')
        if body.get('size', 0) > 20 * 1024 * 1024:
            attachments.append(Item(f'message/{mid}/attachment/{part.get("partId", "0")}',
                'oversized.json', 'application/json', encoded(part),
                text='Attachment exceeds 20 MiB; metadata only.', warning='attachment_too_large'))
            continue
        if body.get('attachmentId'):
            body = request(client, 'GET', root + '/attachments/' + quote(body['attachmentId'], safe=''), headers=headers)
        data = decode64(body.get('data', ''))
        if filename and data:
            attachments.append(Item(f'message/{mid}/attachment/{part.get("partId", "0")}',
                PurePosixPath(filename).name, part.get('mimeType', 'application/octet-stream'),
                data, parse_original=True))
        elif data and part.get('mimeType') in ('text/plain', 'text/html'):
            value = data.decode('utf-8', errors='replace')
            if part['mimeType'] == 'text/html':
                value = html.unescape(re.sub('<[^>]+>', ' ', value))
            text.append(value)
    wanted = {'subject', 'from', 'to', 'date', 'message-id', 'in-reply-to', 'references'}
    heading = '\n'.join(f"{h['name']}: {h['value']}" for h in payload.get('headers', []) if h['name'].lower() in wanted)
    record = Item(f'message/{mid}', f'gmail-{mid}.json', 'application/json', encoded(message),
                  text=f"Gmail message {mid}; thread {message.get('threadId', '')}\n{heading}\n\n" + '\n'.join(text or [message.get('snippet', '')]))
    return [record, *attachments], False


def gmail_batch(client, credentials, config, cursor):
    headers = {'Authorization': 'Bearer ' + credentials['access_token']}
    base = 'https://gmail.googleapis.com/gmail/v1/users/me'
    cursor = dict(cursor)
    labels = config['gmail_label_ids']
    removed = []
    if not cursor.get('history_id'):
        if not cursor.get('baseline'):
            cursor['baseline'] = request(client, 'GET', base + '/profile', headers=headers)['historyId']
        params = {'maxResults': 25, 'labelIds': labels}
        if cursor.get('page'):
            params['pageToken'] = cursor['page']
        response = request(client, 'GET', base + '/messages', headers=headers, params=params)
        ids = [m['id'] for m in response.get('messages', [])]
        next_page = response.get('nextPageToken')
        next_cursor = {**cursor, 'page': next_page} if next_page else {'history_id': cursor['baseline']}
        more = bool(next_page)
    else:
        params = {'startHistoryId': cursor['history_id'], 'maxResults': 25}
        if cursor.get('page'):
            params['pageToken'] = cursor['page']
        try:
            response = request(client, 'GET', base + '/history', headers=headers, params=params)
        except ProviderError as exc:
            if exc.provider_status == 404:
                # Force a fresh snapshot. Existing items are reconciled by the worker at rescan completion.
                return Batch([], [], {'reset': True}, True)
            raise
        ids = set()
        for history in response.get('history', []):
            ids.update(m['id'] for m in history.get('messages', []))
            for field in ('messagesAdded', 'labelsAdded', 'labelsRemoved'):
                ids.update(m['message']['id'] for m in history.get(field, []))
            removed.extend('message/' + m['message']['id'] for m in history.get('messagesDeleted', []))
        ids = sorted(ids)
        next_page = response.get('nextPageToken')
        next_cursor = {**cursor, 'page': next_page} if next_page else {'history_id': response['historyId']}
        more = bool(next_page)
    items = []
    for mid in ids:
        if 'message/' + mid in removed:
            continue
        try:
            found, excluded = gmail_items(client, headers, mid, labels)
        except ProviderError as exc:
            if exc.provider_status == 404:
                removed.append('message/' + mid)
                continue
            raise
        items.extend(found)
        if excluded:
            removed.append('message/' + mid)
    return Batch(items, [], next_cursor, more, tuple(removed))


DRIVE_EXPORT = {
    'application/vnd.google-apps.document': ('text/plain', '.txt'),
    'application/vnd.google-apps.spreadsheet': ('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.xlsx'),
    'application/vnd.google-apps.presentation': ('application/pdf', '.pdf'),
}


def spreadsheet_text(raw):
    from openpyxl import load_workbook
    # XLSX is a zip container: bound expansion as well as downloaded bytes.
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('Expanded workbook too large')
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    lines = ['Spreadsheet source; formulas are preserved, not evaluated.']
    count = 0
    try:
        for sheet in book:
            lines.append('Worksheet: ' + sheet.title)
            for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                count += len(row)
                if count > 100000:
                    raise ValueError('Workbook exceeds 100000 cells')
                if any(v is not None for v in row):
                    lines.append(str(number) + ': ' + ' | '.join('' if v is None else str(v) for v in row))
    finally:
        book.close()
    result = '\n'.join(lines)
    if len(result.encode()) > 20 * 1024 * 1024:
        raise ValueError('Workbook text too large')
    return result


def drive_item(client, headers, file):
    fid, mime = file['id'], file['mimeType']
    if mime == 'application/vnd.google-apps.folder':
        return None
    base = 'https://www.googleapis.com/drive/v3/files/' + quote(fid, safe='')
    filename = PurePosixPath(file['name']).name
    if int(file.get('size') or 0) > 20 * 1024 * 1024:
        return Item('file/' + fid, 'metadata.json', 'application/json', encoded(file),
                    text=f"Drive file {filename}: exceeds 20 MiB; metadata only.", warning='file_too_large')
    if mime in DRIVE_EXPORT:
        mime, ext = DRIVE_EXPORT[mime]
        raw = request(client, 'GET', base + '/export', headers=headers, params={'mimeType': mime}, binary=True)
        filename += ext
    elif mime.startswith('application/vnd.google-apps.'):
        return Item('file/' + fid, 'metadata.json', 'application/json', encoded(file),
                    text=f'Unsupported Google native file: {filename}', warning='unsupported_native_file')
    else:
        raw = request(client, 'GET', base, headers=headers, params={'alt': 'media', 'supportsAllDrives': 'true'}, binary=True)
    return Item('file/' + fid, filename, mime, raw, parse_original=True)


def drive_batch(client, credentials, config, cursor):
    headers = {'Authorization': 'Bearer ' + credentials['access_token']}
    base = 'https://www.googleapis.com/drive/v3'
    cursor = dict(cursor)
    fields = 'id,name,mimeType,modifiedTime,size,trashed,parents'
    folder = config.get('drive_folder_id')
    if not cursor.get('change_token'):
        if not cursor.get('baseline'):
            cursor['baseline'] = request(client, 'GET', base + '/changes/startPageToken', headers=headers,
                params={'supportsAllDrives': 'true'})['startPageToken']
        query = 'trashed = false' + (f" and '{folder}' in parents" if folder else '')
        params = {'pageSize': 25, 'q': query, 'fields': f'nextPageToken,files({fields})',
                  'supportsAllDrives': 'true', 'includeItemsFromAllDrives': 'true'}
        if cursor.get('page'):
            params['pageToken'] = cursor['page']
        response = request(client, 'GET', base + '/files', headers=headers, params=params)
        changes = [{'file': f, 'fileId': f['id']} for f in response.get('files', [])]
        page = response.get('nextPageToken')
        next_cursor = {**cursor, 'page': page} if page else {'change_token': cursor['baseline']}
    else:
        response = request(client, 'GET', base + '/changes', headers=headers, params={
            'pageToken': cursor.get('page') or cursor['change_token'], 'pageSize': 25,
            'supportsAllDrives': 'true', 'includeItemsFromAllDrives': 'true',
            'fields': f'nextPageToken,newStartPageToken,changes(fileId,removed,file({fields}))'})
        changes = response.get('changes', [])
        page = response.get('nextPageToken')
        next_cursor = {**cursor, 'page': page} if page else {'change_token': response['newStartPageToken']}
    items, removed = [], []
    for change in changes:
        file = change.get('file')
        if change.get('removed') or not file or file.get('trashed') or (folder and folder not in file.get('parents', [])):
            removed.append('file/' + change['fileId'])
            continue
        try:
            item = drive_item(client, headers, file)
        except ProviderError as exc:
            if exc.provider_status == 404:
                removed.append('file/' + change['fileId'])
                continue
            raise
        if item:
            items.append(item)
    return Batch(items, removed, next_cursor, bool(page))


def ramp_batch(client, credentials, config, cursor):
    resource = cursor.get('resource', 'bills')
    path = '/developer/v1/' + resource
    host = RAMP_HOSTS[config['environment']]
    params = {'page_size': 100}
    if cursor.get('next'):
        nxt = cursor['next']
        if nxt.startswith(('https:', 'http:', '/')):
            parsed = urlparse(nxt)
            if (parsed.netloc and parsed.netloc != urlparse(host).netloc) or parsed.path != path or (parsed.scheme and parsed.scheme != 'https'):
                raise ConnectorError('Unsafe Ramp pagination URL')
            params = {k: v[-1] for k, v in parse_qs(parsed.query).items()}
        else:
            params['start'] = nxt
    response = request(client, 'GET', host + path, params=params,
                       headers={'Authorization': 'Bearer ' + credentials['access_token']})
    items = []
    for row in response['data']:
        money = row.get('amount')
        basis = 'amount'
        if not isinstance(money, dict):
            basis = 'entity_amount' if row.get('entity_amount') else 'merchant_amount'
            money = row.get(basis)
        normalized = {'id': row['id'], 'provider_id': row['id'], 'status': row.get('status') or row.get('state'),
                      'memo': row.get('memo'), 'invoice_number': row.get('invoice_number'),
                      'vendor_id': (row.get('vendor') or {}).get('id'),
                      'purchase_order_id': row.get('purchase_order_id'), 'currency': None,
                      'amount_minor': None, 'amount_basis': basis,
                      'minor_unit_conversion_rate': None}
        # ApiAmount and value/currency objects have documented minor units. Never guess units for legacy floats.
        if isinstance(money, dict):
            normalized.update(amount_minor=money.get('amount', money.get('value')), currency=money.get('currency_code', money.get('currency')),
                              minor_unit_conversion_rate=money.get('minor_unit_conversion_rate'))
        items.append(Item(resource + '/' + row['id'], resource + '.json', 'application/json', encoded(row),
            record=normalized, dataset='ramp_' + resource, currency=normalized['currency'],
            field_types={'amount_minor': 'numeric', 'minor_unit_conversion_rate': 'numeric'}))
    nxt = (response.get('page') or {}).get('next')
    if nxt:
        return Batch(items, [], {'resource': resource, 'next': nxt}, True)
    if resource == 'bills':
        return Batch(items, [], {'resource': 'transactions'}, True)
    return Batch(items, [], {}, False)


def plaid_batch(client, credentials, config, cursor):
    original = cursor.get('cursor')
    # Collect an entire update before publishing; mutation errors restart from the original cursor.
    for attempt in range(3):
        current, added, removed = original, {}, set()
        try:
            for _ in range(100):
                body = {'access_token': credentials['access_token'], 'count': 500}
                if current:
                    body['cursor'] = current
                response = plaid(client, '/transactions/sync', config['environment'], **body)
                for row in response.get('added', []) + response.get('modified', []):
                    added[row['transaction_id']] = row
                    removed.discard(row['transaction_id'])
                for row in response.get('removed', []):
                    removed.add(row['transaction_id'])
                    added.pop(row['transaction_id'], None)
                current = response['next_cursor']
                if not response['has_more']:
                    break
            else:
                raise ConnectorError('Plaid update exceeds 50000 changes; cursor not advanced')
            break
        except ProviderError as exc:
            if exc.code != 'TRANSACTIONS_SYNC_MUTATION_DURING_PAGINATION' or attempt == 2:
                raise
    items = []
    for row in added.values():
        currency = row.get('iso_currency_code') or row.get('unofficial_currency_code')
        normalized = {k: row.get(k) for k in ('account_id', 'name', 'merchant_name', 'date', 'authorized_date', 'pending', 'pending_transaction_id')}
        normalized.update(id=row['transaction_id'], provider_id=row['transaction_id'],
                          amount_major=str(row['amount']), currency=currency)
        items.append(Item('transaction/' + row['transaction_id'], 'transaction.json', 'application/json', encoded(row),
            record=normalized, dataset='plaid_transactions', currency=currency,
            field_types={'amount_major': 'numeric', 'date': 'date', 'authorized_date': 'date'}))
    return Batch(items, ['transaction/' + rid for rid in sorted(removed)], {'cursor': current})


FETCH = {'gmail': gmail_batch, 'drive': drive_batch, 'ramp': ramp_batch, 'plaid': plaid_batch}
